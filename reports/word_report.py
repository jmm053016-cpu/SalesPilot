"""将统一报告模型渲染为可下载的 Word 文档。"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from models.analysis import MetricValue
from models.report import ArtifactRef, ReportModel

WORD_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _format_value(value: Any) -> str:
    if isinstance(value, MetricValue):
        return str(value.value) if value.value is not None else (value.reason or "不可用")
    return "不可用" if value is None else str(value)


def _metric_label(name: str) -> str:
    return {"sales": "销售额", "net_sales": "净销售额", "orders": "订单数", "quantity": "销量", "average_order_value": "客单价", "gross_profit": "毛利", "gross_margin": "毛利率"}.get(name, name)


def _add_metric_table(document: Document, report: ReportModel) -> None:
    document.add_heading("核心指标", level=1)
    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("指标", "数值", "口径/说明")):
        cell.text = text
    if report.analysis is None:
        row = table.add_row().cells
        row[0].text, row[1].text, row[2].text = "分析结果", "不可用", "缺少结构化分析结果"
        return
    for name, metric in report.analysis.metrics.items():
        row = table.add_row().cells
        row[0].text = _metric_label(name)
        row[1].text = _format_value(metric)
        row[2].text = metric.unit if metric.value is not None else (metric.reason or "不可用")


def _add_rankings(document: Document, report: ReportModel) -> None:
    document.add_heading("品类/产品/区域表现", level=1)
    rankings = report.analysis.rankings if report.analysis else []
    if not rankings:
        document.add_paragraph("暂无可用排名数据。")
        return
    for ranking in rankings:
        document.add_paragraph(f"{ranking.dimension}（指标：{ranking.metric}）")
        keys: list[str] = []
        for item in ranking.items:
            for key in item:
                if key not in keys:
                    keys.append(key)
        table = document.add_table(rows=1, cols=max(1, len(keys)))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, key in zip(table.rows[0].cells, keys):
            cell.text = key
        for item in ranking.items:
            row = table.add_row().cells
            for cell, key in zip(row, keys):
                cell.text = _format_value(item.get(key))


def _add_analysis_sections(document: Document, report: ReportModel) -> None:
    document.add_heading("期间变化", level=1)
    comparisons = report.analysis.comparisons if report.analysis else []
    if not comparisons:
        document.add_paragraph("暂无可用期间对比数据。")
    for comparison in comparisons:
        document.add_paragraph(f"{comparison.dimension}：当前 {_format_value(comparison.current)}，对比期 {_format_value(comparison.previous)}，变化 {_format_value(comparison.change)}。")
    _add_rankings(document, report)
    document.add_heading("异常", level=1)
    anomalies = report.analysis.anomalies if report.analysis else []
    if not anomalies:
        document.add_paragraph("未发现可报告的异常。")
    for anomaly in anomalies:
        document.add_paragraph(f"{anomaly.dimension}：当前值 {_format_value(anomaly.current_value)}，严重程度 {anomaly.severity}。{anomaly.reason or ''}", style="List Bullet")


def _add_narrative(document: Document, report: ReportModel) -> None:
    document.add_heading("经营解读", level=1)
    narrative = report.narrative
    if narrative is None:
        document.add_paragraph("经营解读暂不可用，报告保留确定性分析结果。")
        return
    document.add_paragraph(narrative.summary)
    if narrative.observations:
        document.add_paragraph("观察")
        for item in narrative.observations:
            document.add_paragraph(item.text, style="List Bullet")
    document.add_heading("建议", level=2)
    for item in narrative.recommendations:
        document.add_paragraph(item.text, style="List Bullet")
    if not narrative.recommendations:
        document.add_paragraph("暂无建议。")
    if narrative.hypotheses:
        document.add_heading("待验证假设", level=2)
        for item in narrative.hypotheses:
            document.add_paragraph(item, style="List Bullet")


def export_word(report: ReportModel, output_dir: Path) -> ArtifactRef:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"{report.report_id}.docx"
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    document.styles["Normal"].font.name = "Microsoft YaHei"
    document.styles["Normal"].font.size = Pt(10.5)
    title = report.metadata.title if report.metadata else "销售经营分析报告"
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(18)
    document.add_paragraph(f"分析期间：{report.period.start} 至 {report.period.end}")
    document.add_heading("数据范围与质量", level=1)
    document.add_paragraph(f"数据集：{report.dataset_id}")
    if report.metadata and report.metadata.warnings:
        for warning in report.metadata.warnings:
            document.add_paragraph(warning, style="List Bullet")
    else:
        document.add_paragraph("未记录数据质量警告。")
    _add_metric_table(document, report)
    _add_analysis_sections(document, report)
    _add_narrative(document, report)
    document.add_heading("指标口径附录", level=1)
    document.add_paragraph("报告中的数值均来自 ReportModel 中的确定性分析结果；不可用指标保留其原因，不进行估算。")
    document.save(path)
    return ArtifactRef(artifact_id=f"{report.report_id}_word", kind="report", mime_type=WORD_MIME, display_name=path.name, storage_key=str(path))
