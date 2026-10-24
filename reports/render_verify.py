"""校验文档渲染结果，避免导出文件损坏或内容缺失。"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader
from pydantic import BaseModel, Field


class ReportVerificationResult(BaseModel):
    valid: bool
    errors: list[str] = Field(default_factory=list)
    warnings: list[str] = Field(default_factory=list)
    extracted_text: str = ""


def _word_text(path: Path) -> str:
    document = Document(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join(paragraphs + cells)


def _pdf_text(path: Path) -> str:
    return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)


def verify_report_artifact(ref: Any) -> ReportVerificationResult:
    """Validate a generated report's existence, format, and required sections."""
    path = Path(ref.storage_key)
    errors: list[str] = []
    warnings: list[str] = []
    if not path.exists():
        return ReportVerificationResult(valid=False, errors=[f"artifact does not exist: {path}"])
    if not path.is_file() or path.stat().st_size == 0:
        return ReportVerificationResult(valid=False, errors=[f"artifact is empty or not a file: {path}"])
    expected = {"application/pdf": ".pdf", "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx"}.get(ref.mime_type)
    if expected is None:
        errors.append(f"unsupported artifact MIME type: {ref.mime_type}")
    elif path.suffix.lower() != expected:
        errors.append(f"file extension does not match MIME type: {path.suffix}")
    text = ""
    if not errors:
        try:
            text = _pdf_text(path) if ref.mime_type == "application/pdf" else _word_text(path)
        except Exception as exc:
            errors.append(f"invalid report structure: {exc}")
    required = ("核心指标", "销售额", "建议")
    for section in required:
        if text and section not in text:
            errors.append(f"missing required text: {section}")
    return ReportVerificationResult(valid=not errors, errors=errors, warnings=warnings, extracted_text=text)
