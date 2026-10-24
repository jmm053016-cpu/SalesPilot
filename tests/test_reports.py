from datetime import date
from decimal import Decimal
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from models.analysis import (
    AnalysisRequest,
    AnalysisResult,
    Evidence,
    MetricValue,
    RankingResult,
)
from models.report import NarrativeResult, ReportModel
from models.sales import DateRange
from reports.word_report import export_word
from reports.render_verify import verify_report_artifact


PERIOD = DateRange(start=date(2024, 1, 1), end=date(2024, 1, 31))


def _report(*, with_cost: bool = True, long_table: bool = False) -> ReportModel:
    request = AnalysisRequest(
        dataset_id="dataset_reports",
        current_start=PERIOD.start,
        current_end=PERIOD.end,
        comparison="none",
        group_by=["category", "region"],
    )
    metrics = {
        "sales": MetricValue(value=Decimal("12345.67"), unit="USD"),
        "net_sales": MetricValue(value=Decimal("12000.67"), unit="USD"),
        "orders": MetricValue(value=Decimal("42"), unit="count"),
    }
    evidence = [
        Evidence(
            evidence_id="ev_sales",
            kind="metric",
            label="销售额",
            value=Decimal("12345.67"),
            unit="USD",
            period=PERIOD,
        ),
        Evidence(
            evidence_id="ev_profit",
            kind="metric",
            label="毛利",
            value=Decimal("3456.78"),
            unit="USD",
            period=PERIOD,
        ),
    ]
    if with_cost:
        metrics["gross_profit"] = MetricValue(value=Decimal("3456.78"), unit="USD")
        metrics["gross_margin"] = MetricValue(value=Decimal("0.2880"), unit="ratio")
    else:
        metrics["gross_profit"] = MetricValue(value=None, unit="USD", reason="缺少成本字段，无法计算毛利")
        metrics["gross_margin"] = MetricValue(value=None, unit="ratio", reason="缺少成本字段，无法计算毛利率")

    items = [
        {
            "name": f"品类-{index:03d}",
            "metric": MetricValue(value=Decimal(index * 100), unit="USD"),
        }
        for index in range(1, 61 if long_table else 4)
    ]
    analysis = AnalysisResult(
        analysis_id="analysis_reports",
        dataset_id="dataset_reports",
        request=request,
        metrics=metrics,
        rankings=[RankingResult(dimension="category", metric="sales", period=PERIOD, items=items)],
        evidence=evidence,
    )
    return ReportModel(
        report_id="report_reports",
        analysis_id=analysis.analysis_id,
        dataset_id=analysis.dataset_id,
        period=PERIOD,
        analysis=analysis,
        narrative=NarrativeResult(
            summary="销售额保持稳定，核心品类贡献了主要收入。",
            observations=[{"text": "销售额为 12345.67 USD。", "evidence_ids": ["ev_sales"]}],
            recommendations=[{"text": "继续关注核心品类表现。", "evidence_ids": ["ev_sales"]}],
            hypotheses=["需要进一步验证区域差异。"],
            status="generated",
        ),
    )


def _artifact_path(ref) -> Path:
    return Path(ref.storage_key)


def test_export_word_creates_openable_report_with_core_metrics_and_sections(tmp_path):
    ref = export_word(_report(), tmp_path)

    assert ref.mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    path = _artifact_path(ref)
    assert path.exists() and path.stat().st_size > 0

    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    text += "\n" + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "核心指标" in text
    assert "销售额" in text
    assert "12345.67" in text
    assert "毛利" in text
    assert "3456.78" in text
    assert "异常" in text
    assert "建议" in text


def test_export_pdf_creates_openable_report_with_core_metrics_and_sections(tmp_path):
    from reports.pdf_report import export_pdf

    ref = export_pdf(_report(), tmp_path)

    assert ref.mime_type == "application/pdf"
    path = _artifact_path(ref)
    assert path.exists() and path.stat().st_size > 0

    reader = PdfReader(str(path))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "核心指标" in text
    assert "销售额" in text
    assert "12345.67" in text
    assert "毛利" in text
    assert "3456.78" in text
    assert "建议" in text


def test_exports_explicitly_explain_unavailable_profit_when_cost_is_missing(tmp_path):
    from reports.pdf_report import export_pdf

    word_ref = export_word(_report(with_cost=False), tmp_path)
    pdf_ref = export_pdf(_report(with_cost=False), tmp_path)

    word = Document(_artifact_path(word_ref))
    word_text = "\n".join(p.text for p in word.paragraphs)
    word_text += "\n" + "\n".join(cell.text for table in word.tables for row in table.rows for cell in row.cells)
    pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(str(_artifact_path(pdf_ref))).pages)
    expected = "缺少成本字段，无法计算毛利"
    assert expected in word_text
    assert expected in pdf_text


def test_exports_handle_long_tables_and_chinese_text(tmp_path):
    from reports.pdf_report import export_pdf

    word_ref = export_word(_report(long_table=True), tmp_path)
    pdf_ref = export_pdf(_report(long_table=True), tmp_path)

    assert len(Document(_artifact_path(word_ref)).tables) >= 1
    assert len(PdfReader(str(_artifact_path(pdf_ref))).pages) >= 2

def test_verify_report_artifact_validates_word_and_pdf_structure(tmp_path):
    word_ref = export_word(_report(), tmp_path)
    pdf_ref = __import__('reports.pdf_report', fromlist=['export_pdf']).export_pdf(_report(), tmp_path)

    word_result = verify_report_artifact(word_ref)
    pdf_result = verify_report_artifact(pdf_ref)

    assert word_result.valid is True
    assert pdf_result.valid is True
    assert word_result.errors == []
    assert pdf_result.errors == []

def test_verify_report_artifact_rejects_missing_or_malformed_file(tmp_path):
    missing = type('Ref', (), {'storage_key': str(tmp_path / 'missing.pdf'), 'mime_type': 'application/pdf'})()

    result = verify_report_artifact(missing)

    assert result.valid is False
    assert any('does not exist' in error for error in result.errors)
