"""读取 Excel 工作簿并选择目标工作表作为导入数据。"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from models.sales import ImportOptions, RawDataset

from .excel_formulas import ExcelFormulaEvaluator, formula_result_as_text

from .schema import MAX_ROWS, check_source_limits


def load_excel(path: Path, sheet_name: str | int, options: ImportOptions) -> RawDataset:
    path = Path(path)
    check_source_limits(path)
    frame = pd.read_excel(path, sheet_name=sheet_name, engine='openpyxl', dtype=str)
    if len(frame) > MAX_ROWS:
        raise ValueError('row count exceeds 100,000 limit')

    workbook = load_workbook(path, data_only=False, read_only=False)
    try:
        worksheet = workbook.worksheets[sheet_name] if isinstance(sheet_name, int) else workbook[sheet_name]
        evaluator = ExcelFormulaEvaluator(worksheet)
        for row_index in range(2, len(frame) + 2):
            for column_index in range(1, frame.shape[1] + 1):
                cell = worksheet.cell(row=row_index, column=column_index)
                if isinstance(cell.value, str) and cell.value.startswith('='):
                    frame.iat[row_index - 2, column_index - 1] = formula_result_as_text(evaluator.value(cell.coordinate))
    finally:
        workbook.close()
    return RawDataset(frame, options)
